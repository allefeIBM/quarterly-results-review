#!/usr/bin/env python3
"""
update_data.py — Regenera o data.json a partir dos arquivos Q126/Q226/Q326/Q426.xlsx.

Formato das planilhas:
  Linha 1  → cabeçalho: UT15 (KPI) | UT17 (KPI) | Activity Name | Brief/EMB | CQ | NQ | NQ+1 | NQ+2 | Total
  Linha 2+ → dados

Uso:
    python3 update_data.py

O script lê os 4 .xlsx, extrai os KPIs e salva o data.json atualizado.
Em seguida, faz commit e push automático para o GitHub (opcional).
"""

import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).parent

# Planilhas a processar (na ordem Q1→Q4)
XLSX_FILES = [
    ("Q126.xlsx", "Q1"),
    ("Q226.xlsx", "Q2"),
    ("Q326.xlsx", "Q3"),
    ("Q426.xlsx", "Q4"),
]


# ── Helpers de parsing ────────────────────────────────────────────────────────

def col_index_exact(headers, *names):
    """Retorna o índice da primeira coluna cujo cabeçalho (sem sufixo entre parênteses)
    corresponde a um dos nomes fornecidos (case-insensitive)."""
    for name in names:
        nl = name.upper()
        for i, h in enumerate(headers):
            hc = str(h).strip().upper().split("(")[0].strip() if h else ""
            if hc == nl:
                return i
    return None


def col_index_contains(headers, *fragments):
    """Retorna o índice da primeira coluna cujo cabeçalho contém um dos fragmentos
    (case-insensitive). Usado como fallback."""
    for frag in fragments:
        fl = frag.lower()
        for i, h in enumerate(headers):
            if h and fl in str(h).strip().lower():
                return i
    return None


def parse_num(value):
    """Converte valor de célula para float ou None."""
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# ── Extração principal ────────────────────────────────────────────────────────

def extract_all_events():
    events = {}

    for filename, quarter in XLSX_FILES:
        path = BASE_DIR / filename
        if not path.exists():
            print(f"  ⚠  {filename} não encontrado — ignorando.")
            continue

        print(f"  → {filename}  (quarter={quarter})")
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:
            print(f"     ⚠  Erro ao abrir: {e}")
            continue

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Linha 1 é sempre o cabeçalho
            raw_headers = rows[0]
            headers = [str(h).strip() if h else "" for h in raw_headers]

            # Mapear colunas pelo cabeçalho
            icol = {
                "ut15":  col_index_exact(headers, "UT15"),
                "ut17":  col_index_exact(headers, "UT17"),
                "name":  col_index_exact(headers, "ACTIVITY NAME")
                         or col_index_contains(headers, "activity name", "event name"),
                "brief": col_index_exact(headers, "BRIEF", "EMB", "EMB #", "EMB#"),
                "cq":    col_index_exact(headers, "CQ"),
                "nq":    col_index_exact(headers, "NQ"),
                "nq1":   col_index_exact(headers, "NQ+1"),
                "nq2":   col_index_exact(headers, "NQ+2"),
                "total": col_index_exact(headers, "TOTAL", "TOTALS", "TOTAL SQL"),
            }

            if icol["name"] is None or icol["brief"] is None:
                print(f"     ⚠  Aba '{sheet_name}' sem colunas obrigatórias — ignorando.")
                continue

            row_count = 0
            for row in rows[1:]:
                name_val  = row[icol["name"]]  if icol["name"]  is not None else None
                brief_val = row[icol["brief"]] if icol["brief"] is not None else None

                name_str  = str(name_val).strip()  if name_val  else ""
                brief_str = str(brief_val).strip() if brief_val else ""

                if not brief_str:
                    continue

                cq    = parse_num(row[icol["cq"]])    if icol.get("cq")    is not None else None
                nq    = parse_num(row[icol["nq"]])    if icol.get("nq")    is not None else None
                nq1   = parse_num(row[icol["nq1"]])   if icol.get("nq1")   is not None else None
                nq2   = parse_num(row[icol["nq2"]])   if icol.get("nq2")   is not None else None
                total = parse_num(row[icol["total"]]) if icol.get("total") is not None else None

                if all(x is None for x in [cq, nq, nq1, nq2, total]):
                    continue

                ut15 = str(row[icol["ut15"]]).strip() if icol.get("ut15") is not None and row[icol["ut15"]] else ""
                ut17 = str(row[icol["ut17"]]).strip() if icol.get("ut17") is not None and row[icol["ut17"]] else ""

                kpi_row = {
                    "event_name":  name_str,
                    "brief":       brief_str,
                    "ut15":        ut15,
                    "ut17":        ut17,
                    "cq":          cq,
                    "nq":          nq,
                    "nq1":         nq1,
                    "nq2":         nq2,
                    "total":       total,
                    "quarter":     quarter,
                    "source_file": filename,
                }

                key = brief_str.upper()
                if key not in events:
                    events[key] = {
                        "event_name": name_str,
                        "brief":      brief_str,
                        "kpi_rows":   [],
                    }
                events[key]["kpi_rows"].append(kpi_row)
                if name_str:
                    events[key]["event_name"] = name_str
                row_count += 1

            print(f"     ✓ aba '{sheet_name}' → {row_count} linha(s) carregada(s)")

    return events


# ── Escrita do data.json ──────────────────────────────────────────────────────

def write_data_json(events):
    output = {
        "generated": datetime.now().isoformat()[:16],
        "events":    events,
    }
    out_path = BASE_DIR / "data.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, separators=(",", ":"))
    size_kb = out_path.stat().st_size / 1024
    print(f"\n✅ data.json gerado — {len(events)} eventos · {size_kb:.1f} KB")
    return out_path


# ── Git commit + push ─────────────────────────────────────────────────────────

def git_push(out_path):
    print()
    answer = input("Deseja fazer commit e push para o GitHub agora? [s/N] ").strip().lower()
    if answer not in ("s", "sim", "y", "yes"):
        print("ℹ  Push ignorado. Rode 'git add data.json && git commit -m \"data: update\" && git push' manualmente.")
        return

    try:
        subprocess.run(["git", "add", str(out_path)], check=True, cwd=BASE_DIR)
        msg = f"data: update KPI projections — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        subprocess.run(["git", "commit", "-m", msg], check=True, cwd=BASE_DIR)
        subprocess.run(["git", "push"], check=True, cwd=BASE_DIR)
        print("🚀 Push concluído! A página será atualizada em ~1 minuto.")
    except subprocess.CalledProcessError as e:
        print(f"⚠  Erro no git: {e}")
        print("   Rode manualmente: git add data.json && git commit -m 'data: update' && git push")


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print("=" * 60)
    print("  Quarterly Results Review — Atualizar data.json")
    print("=" * 60)
    print()

    events = extract_all_events()
    out_path = write_data_json(events)
    git_push(out_path)
